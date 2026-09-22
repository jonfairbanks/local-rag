from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from pptx import Presentation

from utils.llama_index import load_documents


class OfficeDocumentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)

    def write_docx(self):
        with ZipFile(self.directory / "warehouse.docx", "w") as archive:
            archive.writestr("[Content_Types].xml", '''
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>''')
            archive.writestr("_rels/.rels", '''
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>''')
            archive.writestr("word/document.xml", '''
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body><w:p><w:r><w:t>Warehouse inventory code: CEDAR-582.</w:t></w:r></w:p></w:body>
</w:document>''')

    def write_pptx(self):
        slides = Presentation()
        slide = slides.slides.add_slide(slides.slide_layouts[1])
        slide.shapes.title.text = "Planetarium"
        slide.placeholders[1].text = "Project code: ORBIT-426."
        slides.save(self.directory / "planetarium.pptx")

    def loaded_text(self):
        return "\n".join(document.text for document in load_documents(str(self.directory)))

    def test_docx_text_is_loaded(self):
        self.write_docx()
        self.assertIn("CEDAR-582", self.loaded_text())

    def test_pptx_text_is_loaded(self):
        self.write_pptx()
        self.assertIn("ORBIT-426", self.loaded_text())

    def test_mixed_office_uploads_keep_both_documents(self):
        self.write_docx()
        self.write_pptx()
        text = self.loaded_text()
        self.assertIn("CEDAR-582", text)
        self.assertIn("ORBIT-426", text)


if __name__ == "__main__":
    unittest.main()
